"""
generar_ficha.py
================
Genera la ficha informativa trimestral (.txt y .docx) a partir de los
indicadores de calcular_indicadores.py.

Todos los textos y cifras se arman una sola vez en `preparar()` y después se
escriben en los dos formatos, así el .txt y el .docx siempre dicen lo mismo.

Uso:
    python generar_ficha.py --periodo 2026_T2
"""

import argparse
import re
import textwrap
from pathlib import Path
from types import SimpleNamespace

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from calcular_indicadores import JALISCO, REGION, SECTORES, calcular

REGLA = "─" * 60
SERIES = "Series originales (sin ajuste estacional), periodicidad trimestral."
NOTA_SEXO = "Las estimaciones por sexo tienen mayor error de muestreo; cautela con variaciones pequeñas."
ORDINAL = {1: "1°", 2: "2°", 3: "3°", 4: "4°"}
NOMBRES = {
    "td":           "Tasa de Desocupación (TD)",
    "tpea":         "Tasa de Participación Económica (TPEA)",
    "til1":         "Tasa de Informalidad Laboral (TIL1)",
    "subocupacion": "Tasa de Subocupación",
}
CORTOS = {"td": "TD", "tpea": "TPEA", "til1": "TIL1", "subocupacion": "Subocupación"}
DECIMALES = {"td": 2, "tpea": 1, "til1": 1, "subocupacion": 1}


# ── Formato de números ────────────────────────────────────────────────────────
# Las diferencias se calculan con los valores ya redondeados para que cada
# tabla sea consistente a simple vista (lo que se lee es lo que se resta).

def r(v, k):
    return round(v, DECIMALES[k])

def pct(v, k):
    return f"{r(v, k):.{DECIMALES[k]}f}%"

def dif(a, b, k):
    return round(r(a, k) - r(b, k), DECIMALES[k])

def signo(v, dec=2):
    return f"{'−' if v < 0 else '+'}{abs(v):.{dec}f}"

def miles(n):
    return f"{n:,}"

def con_signo(n):
    """Entero con signo y separador de miles: +12,345 / −3,210."""
    return f"{'−' if n < 0 else '+'}{abs(n):,}"

def cambio_pct(actual, base):
    return signo((actual - base) / base * 100, 1) + "%"

def var_personas(actual, base):
    """'+12,345 (+0.3%)': cambio en personas y en porcentaje."""
    return f"{con_signo(actual - base)} ({cambio_pct(actual, base)})"

def etiqueta(periodo):
    anio, trim = periodo.split("_T")
    return f"T{trim} {anio}"


# ── Contenido de la ficha ─────────────────────────────────────────────────────

def preparar(act: dict, ant: dict, dic: dict | None = None, prev: dict | None = None) -> SimpleNamespace:
    """Arma todas las cifras y tablas de la ficha.
    dic = T4 del año anterior (cierre de diciembre); prev = trimestre anterior.
    Sin ellos (p. ej. desde el dashboard) se omite la sección de población ocupada."""
    anio, trim = map(int, re.fullmatch(r"(\d{4})_T(\d)", act["periodo"]).groups())
    c = SimpleNamespace()
    c.anio, c.trim = anio, trim
    c.act, c.ant = f"T{trim} {anio}", f"T{trim} {anio - 1}"
    c.titulo = f"Tasa de Desocupación — Jalisco. {ORDINAL[trim]} Trimestre {anio}"

    J, Ja = act["jalisco"]["total"], ant["jalisco"]["total"]
    N, Na = act["nacional"], ant["nacional"]
    c.J, c.Ja, c.N = J, Ja, N
    c.sexo = {s: (act["jalisco"][s], ant["jalisco"][s]) for s in ("hombres", "mujeres")}

    # ── Análisis anual ────────────────────────────────────────────────────
    var = dif(J["td"], Ja["td"], "td")
    dif_nac = dif(J["td"], N["td"], "td")
    c.t_nacional = (
        ["Período", "Jalisco", "Nacional", "Diferencia (pp)"],
        [
            [c.ant, pct(Ja["td"], "td"), pct(Na["td"], "td"), signo(dif(Ja["td"], Na["td"], "td"))],
            [c.act, pct(J["td"], "td"), pct(N["td"], "td"), signo(dif_nac)],
            ["Var. pp", signo(var), signo(dif(N["td"], Na["td"], "td")), "—"],
        ],
    )

    c.x_anual = [
        ("La TD es el porcentaje de la población económicamente activa que no tiene trabajo y lo "
         f"busca; una variación negativa es favorable. En el {c.act} Jalisco registró ", 0),
        (pct(J["td"], "td"), 1), (" (", 0), (f"{signo(var)} pp", 1), (f" frente al {c.ant}) y se ubicó ", 0),
        (f"{abs(dif_nac):.2f} pp {'por debajo' if dif_nac < 0 else 'por encima'}" if dif_nac
         else "en el mismo nivel", 1),
        (f" del promedio nacional ({pct(N['td'], 'td')}).", 0),
    ]

    # ── Posicionamiento nacional ──────────────────────────────────────────
    ents = sorted(act["entidades"], key=lambda e: e["td"])
    n = len(ents)
    rank = next(i for i, e in enumerate(ents, 1) if e["cve"] == JALISCO)
    c.rank, c.n = rank, n
    if rank <= 8:
        zona = "entre las entidades con menor desocupación"
    elif rank > n - 8:
        zona = "entre las entidades con mayor desocupación"
    else:
        zona = "en la parte media de la tabla"
    til_bajas = sum(r(e["til1"], "til1") for e in ents[:5]) / 5
    c.x_ranking = [
        (f"Las {n} entidades se ordenan de menor a mayor TD (lugar 1 = menor desocupación). Jalisco ocupa el ", 0),
        (f"lugar {rank} de {n}", 1), (f", {zona}. Una TD baja no siempre indica un mejor mercado laboral: "
         f"las cinco entidades con menor TD tienen una informalidad promedio de {til_bajas:.1f}% "
         f"(nacional: {pct(N['til1'], 'til1')}).", 0),
    ]
    filas = lambda lista, inicio: [[str(i), e["nombre"], pct(e["td"], "td")]
                                    for i, e in enumerate(lista, inicio)]
    c.t_menor = filas(ents[:5], 1)
    if rank > 5:
        c.t_menor += [["…", "…", "…"], [str(rank), "Jalisco ←", pct(J["td"], "td")]]
    c.t_mayor = [[str(i), e["nombre"], pct(e["td"], "td")]
                 for i, e in enumerate(ents[::-1][:5], 1)]

    # ── Comparativo regional ──────────────────────────────────────────────
    reg = [e for e in ents if e["cve"] in REGION]
    media = round(sum(r(e["td"], "td") for e in reg) / len(reg), 2)
    rank_reg = next(i for i, e in enumerate(reg, 1) if e["cve"] == JALISCO)
    d_media = round(r(J["td"], "td") - media, 2)
    c.x_region = [
        (f"Compara a Jalisco con las {len(reg)} entidades de la región; la última columna es la distancia "
         "a la media regional (negativo = menor desocupación que la media). Jalisco ocupa el ", 0),
        (f"{rank_reg}° lugar", 1), (" con menor TD y está ", 0),
        ("prácticamente en la media" if abs(d_media) < 0.05
         else f"{abs(d_media):.2f} pp {'por debajo' if d_media < 0 else 'por encima'} de la media", 1),
        (f" ({media:.2f}%).", 0),
    ]
    c.t_region = [
        [str(i), e["nombre"] + (" ←" if e["cve"] == JALISCO else ""), pct(e["td"], "td"),
         signo(round(r(e["td"], "td") - media, 2))]
        for i, e in enumerate(reg, 1)
    ] + [["", "Media regional", f"{media:.2f}%", "—"]]

    # ── Indicadores complementarios ───────────────────────────────────────
    c.t_indicadores = [
        [NOMBRES[k], pct(Ja[k], k), pct(J[k], k),
         signo(dif(J[k], Ja[k], k), DECIMALES[k]), pct(N[k], k)]
        for k in NOMBRES
    ]

    v = {k: dif(J[k], Ja[k], k) for k in NOMBRES}
    lectura = ""
    if v["td"] < 0 and v["tpea"] < 0:
        lectura = (" La baja de la TD coincide con una menor participación: parte de la mejora puede deberse "
                   "a que menos personas buscan empleo, no solo a más contratación.")
    elif v["td"] > 0 and v["tpea"] > 0:
        lectura = " El aumento de la TD coincide con más personas participando o buscando empleo."
    d_til = dif(J["til1"], N["til1"], "til1")
    c.x_indicadores = [
        ("Complementan la lectura de la TD: la TPEA indica qué parte de la población de 15 años y más trabaja "
         "o busca trabajo; la TIL1, qué parte de los ocupados es informal, y la subocupación, qué parte de los "
         f"ocupados necesita trabajar más horas. Frente al {c.ant}: ", 0),
        (", ".join(f"{CORTOS[k]} {signo(v[k], DECIMALES[k])} pp" for k in NOMBRES), 1),
        (f".{lectura} La informalidad se mantiene {abs(d_til):.1f} pp "
         f"{'por debajo' if d_til < 0 else 'por encima'} del promedio nacional.", 0),
    ]

    # ── Indicadores por sexo ──────────────────────────────────────────────
    (H, Ha), (M, Ma) = c.sexo["hombres"], c.sexo["mujeres"]
    b_td, b_tpea = dif(M["td"], H["td"], "td"), dif(M["tpea"], H["tpea"], "tpea")
    c.x_sexo = [
        ("La brecha M−H es la tasa de las mujeres menos la de los hombres (negativo = menor en mujeres). "
         "La TD de las mujeres (", 0), (pct(M["td"], "td"), 1),
        (f") es {'menor' if b_td < 0 else 'mayor' if b_td > 0 else 'igual'} que la de los hombres (", 0),
        (pct(H["td"], "td"), 1), ("), pero su participación es ", 0),
        (f"{abs(b_tpea):.1f} pp {'menor' if b_tpea < 0 else 'mayor'}", 1),
        (f" ({pct(M['tpea'], 'tpea')} vs {pct(H['tpea'], 'tpea')}): la TD de cada grupo debe leerse junto con "
         "su participación.", 0),
    ]
    cambio = lambda x, d: f"{'subió' if x > 0 else 'bajó'} {abs(x):.{d}f} pp" if x else "no cambió"
    c.x_sexo_var = [
        (f"Cambio en puntos porcentuales frente al {c.ant} (positivo = aumento). La TD de las mujeres ", 0),
        (cambio(dif(M["td"], Ma["td"], "td"), 2), 1), (" y la de los hombres ", 0),
        (cambio(dif(H["td"], Ha["td"], "td"), 2), 1), (".", 0),
    ]
    c.t_sexo = [
        [CORTOS[k], pct(J[k], k), pct(H[k], k), pct(M[k], k),
         f"{signo(dif(M[k], H[k], k), DECIMALES[k])} pp"]
        for k in NOMBRES
    ]
    c.t_sexo_var = [
        [CORTOS[k], signo(dif(H[k], Ha[k], k), DECIMALES[k]), signo(dif(M[k], Ma[k], k), DECIMALES[k])]
        for k in NOMBRES
    ]

    # ── Población ocupada ─────────────────────────────────────────────────
    c.t_ocupacion, c.nota_sectores = None, ""
    c.ocupada_resumen = "n/d"
    if dic and prev and all("ocupacion" in i["jalisco"] for i in (act, ant, dic, prev)):
        _preparar_ocupacion(c, act, ant, dic, prev)

    return c


def _preparar_ocupacion(c, act, ant, dic, prev):
    O = lambda ind: ind["jalisco"]["ocupacion"]
    o = O(act)

    # Las tres vistas de comparación. En el T1, el cierre de diciembre y el
    # trimestre anterior son el mismo (T4 del año anterior): una sola columna.
    if dic["periodo"] == prev["periodo"]:
        comparaciones = [dic, ant]
    else:
        comparaciones = [dic, ant, prev]

    filas_def = [("Población ocupada total", lambda x: x["total"]),
                 ("Formalidad", None),
                 ("  Formal", lambda x: x["formal"]),
                 ("  Informal", lambda x: x["informal"]),
                 ("Sexo", None),
                 ("  Hombres", lambda x: x["hombres"]),
                 ("  Mujeres", lambda x: x["mujeres"]),
                 ("Actividad económica", None)]
    filas_def += [(f"  {nombre}", lambda x, k=str(k): x["sectores"][k]) for k, nombre in SECTORES.items()]
    if o["sector_no_especificado"] > 0:
        filas_def.append(("  No especificado", lambda x: x["sector_no_especificado"]))

    encabezados = ["Concepto", c.act] + [f"vs {etiqueta(i['periodo'])}" for i in comparaciones]
    filas, grupos = [], []
    for nombre, valor in filas_def:
        if valor is None:
            grupos.append(len(filas))
            filas.append([nombre] + [""] * (len(encabezados) - 1))
        else:
            filas.append([nombre, miles(valor(o))] + [var_personas(valor(o), valor(O(i))) for i in comparaciones])
    c.t_ocupacion = (encabezados, filas)
    c.grupos_ocupacion = [0] + grupos           # filas en negritas en el Word

    # Nota breve: sectores con muy pocos casos en la muestra (menos del 1% de los ocupados)
    chicos = [nombre for k, nombre in SECTORES.items() if o["sectores"][str(k)] < 0.01 * o["total"]]
    if o["sector_no_especificado"] > 0:
        chicos.append("No especificado")
    if chicos:
        c.nota_sectores = f"Sectores con pocos casos ({', '.join(chicos)}): variaciones poco confiables."

    total, total_ant = o["total"], O(ant)["total"]
    if len(comparaciones) == 3:
        refs = (f"al cierre del año anterior ({etiqueta(dic['periodo'])}), al mismo trimestre del año anterior "
                f"({etiqueta(ant['periodo'])}; la comparación más limpia porque elimina el efecto estacional) "
                f"y al trimestre anterior ({etiqueta(prev['periodo'])})")
    else:
        refs = (f"al cierre del año anterior ({etiqueta(dic['periodo'])}, que también es el trimestre anterior) "
                f"y al mismo trimestre del año anterior ({etiqueta(ant['periodo'])}; la comparación más limpia "
                "porque elimina el efecto estacional)")
    confiables = {nombre: o["sectores"][str(k)] - O(ant)["sectores"][str(k)]
                  for k, nombre in SECTORES.items() if nombre not in chicos}
    orden = sorted(confiables.items(), key=lambda kv: kv[1])
    suben = [f"{n} ({con_signo(v)})" for n, v in orden[::-1][:2] if v > 0]
    bajan = [f"{n} ({con_signo(v)})" for n, v in orden[:2] if v < 0]
    dif_total = total - total_ant
    sectores_txt = ""
    if suben:
        sectores_txt += f" Por actividad, los mayores aumentos anuales fueron en {' y '.join(suben)}"
    if bajan:
        sectores_txt += ("; las mayores caídas, en " if suben else " Por actividad, las mayores caídas anuales fueron en ") \
            + " y ".join(bajan)
    c.x_ocupacion = [
        (f"Personas ocupadas y su cambio respecto {refs}. En el último año la población ocupada ", 0),
        (f"{'aumentó' if dif_total > 0 else 'disminuyó'} en {miles(abs(dif_total))} personas "
         f"({cambio_pct(total, total_ant)})" if dif_total else "no cambió", 1),
        (f".{sectores_txt}." if sectores_txt else ".", 0),
    ]
    c.ocupada_resumen = f"{miles(total)} personas ({cambio_pct(total, total_ant)} anual)"


# ── Salida .txt ───────────────────────────────────────────────────────────────

def _tabla_txt(encabezados, filas):
    anchos = [max(len(str(x)) for x in col) for col in zip(encabezados, *filas)]
    linea = lambda celdas: "  ".join(str(x).ljust(a) for x, a in zip(celdas, anchos)).rstrip()
    return "\n".join([linea(encabezados), linea(["─" * a for a in anchos])]
                     + [linea(f) for f in filas])


def _texto_txt(partes, ancho=95):
    # Los espacios dentro de negritas se protegen para no partir un *...* en dos líneas
    texto = "".join(f"*{t.replace(' ', chr(0))}*" if b else t for t, b in partes)
    return textwrap.fill(texto, ancho, break_on_hyphens=False).replace(chr(0), " ")


def generar_txt(c: SimpleNamespace, ruta):
    s = []
    seccion = lambda titulo, sub="": s.extend(
        ["", REGLA, f"*{titulo}*"] + ([sub] if sub else []) + [REGLA, ""])

    s += [f"*{c.titulo}*", "", f"_*{SERIES}_"]

    seccion("Análisis Anual", "Último trimestre disponible vs mismo trimestre del año anterior.")
    s += [_texto_txt(c.x_anual), "", f"*Comparativo Jalisco vs Nacional — {c.act}:*", _tabla_txt(*c.t_nacional)]

    seccion("Posicionamiento Nacional", f"TD de Jalisco vs entidades federativas — {c.act}")
    s += [_texto_txt(c.x_ranking), "", f"*Top 5 entidades con menor TD ({c.act}):*",
          _tabla_txt(["Lugar", "Estado", "TD (%)"], c.t_menor), "",
          f"*Top 5 entidades con mayor TD ({c.act}):*",
          _tabla_txt(["Lugar", "Estado", "TD (%)"], c.t_mayor)]

    seccion(f"Comparativo Regional — Occidente / Centro-Bajío ({c.act})")
    s += [_texto_txt(c.x_region), "", _tabla_txt(["Lugar", "Estado", "TD (%)", "vs Media regional (pp)"], c.t_region)]

    seccion("Indicadores Complementarios del Mercado Laboral",
            f"Comparativo anual: {c.act} vs {c.ant} — Jalisco")
    s += [_texto_txt(c.x_indicadores), "",
          _tabla_txt(["Indicador", c.ant, c.act, "Var. (pp)", f"Nacional {c.act}"], c.t_indicadores)]

    seccion(f"Indicadores por Sexo — Jalisco {c.act}")
    s += [_texto_txt(c.x_sexo), "", _tabla_txt(["Indicador", "Total", "Hombres", "Mujeres", "Brecha M−H"], c.t_sexo), "",
          f"*Variación anual por sexo ({c.ant} → {c.act}, pp):*", _texto_txt(c.x_sexo_var), "",
          _tabla_txt(["Indicador", "Hombres", "Mujeres"], c.t_sexo_var), "", f"_{NOTA_SEXO}_"]

    if c.t_ocupacion:
        seccion(f"Población Ocupada — Jalisco {c.act}", "Personas y variación (absoluta y %)")
        s += [_texto_txt(c.x_ocupacion), "", _tabla_txt(*c.t_ocupacion)]
        if c.nota_sectores:
            s += ["", f"_{c.nota_sectores}_"]

    Path(ruta).write_text("\n".join(s) + "\n", encoding="utf-8")


# ── Salida .docx ──────────────────────────────────────────────────────────────

FUENTE = "Calibri"
T_NORMAL, T_CHICA, T_TITULO = Pt(10), Pt(9), Pt(13)


def _run(p, texto, size=T_NORMAL, bold=False, italic=False):
    run = p.add_run(texto)
    run.font.name, run.font.size, run.font.bold, run.font.italic = FUENTE, size, bold, italic


def _p(doc, partes=(), size=T_NORMAL, italic=False, antes=3, despues=3, sangria=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(antes), Pt(despues)
    if sangria:
        p.paragraph_format.left_indent = Cm(sangria)
    if isinstance(partes, str):
        partes = [(partes, 0)]
    for texto, bold in partes:
        _run(p, texto, size, bool(bold), italic)
    return p


def _regla(doc):
    p = _p(doc, antes=4, despues=4)
    p.paragraph_format.keep_with_next = True
    borde = OxmlElement("w:pBdr")
    abajo = OxmlElement("w:bottom")
    for k, v in {"val": "single", "sz": "6", "space": "1", "color": "888888"}.items():
        abajo.set(qn(f"w:{k}"), v)
    borde.append(abajo)
    p._p.get_or_add_pPr().append(borde)


def _seccion(doc, titulo, sub=""):
    _regla(doc)
    p = _p(doc, [(titulo, 1)], antes=8, despues=2)
    p.paragraph_format.keep_with_next = True
    if sub:
        _run(p, f"  —  {sub}", T_CHICA)


def _tabla(doc, encabezados, filas, anchos_cm, primera_izq=False, negritas=()):
    # El párrafo anterior (título de la tabla) se mantiene junto a la tabla
    doc.paragraphs[-1].paragraph_format.keep_with_next = True
    t = doc.add_table(rows=1 + len(filas), cols=len(encabezados))
    t.style, t.alignment, t.autofit = "Table Grid", WD_TABLE_ALIGNMENT.LEFT, False
    for j, ancho in enumerate(anchos_cm):
        t.columns[j].width = Cm(ancho)
    ultima = len(filas)
    for i, fila in enumerate([encabezados] + filas):
        for j, valor in enumerate(fila):
            celda = t.cell(i, j)
            celda.width = Cm(anchos_cm[j])
            p = celda.paragraphs[0]
            izq = primera_izq and j == (1 if encabezados[0] == "Lugar" else 0)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if izq else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = i < ultima   # la tabla no se parte entre páginas
            _run(p, str(valor), T_CHICA, bold=(i == 0 or (i - 1) in negritas))
    _p(doc, antes=0, despues=4)


def generar_docx(c: SimpleNamespace, ruta):
    doc = Document()
    for sec in doc.sections:
        sec.top_margin = sec.bottom_margin = Cm(2.5)
        sec.left_margin = sec.right_margin = Cm(3.0)

    _p(doc, [(c.titulo, 1)], size=T_TITULO, antes=0, despues=4)
    _p(doc, SERIES, T_CHICA, italic=True, antes=0, despues=4)

    _seccion(doc, "Análisis Anual", "Último trimestre disponible vs mismo trimestre del año anterior")
    _p(doc, c.x_anual, despues=5)
    _p(doc, [(f"Comparativo Jalisco vs Nacional — {c.act}:", 1)])
    _tabla(doc, *c.t_nacional, [3.2, 3.5, 3.5, 4.5])

    _seccion(doc, "Posicionamiento Nacional", f"TD de Jalisco vs entidades federativas — {c.act}")
    _p(doc, c.x_ranking, despues=5)
    _p(doc, [(f"Top 5 entidades con menor TD ({c.act}):", 1)])
    _tabla(doc, ["Lugar", "Estado", "TD (%)"], c.t_menor, [2.0, 8.5, 4.2], primera_izq=True)
    _p(doc, [(f"Top 5 entidades con mayor TD ({c.act}):", 1)])
    _tabla(doc, ["Lugar", "Estado", "TD (%)"], c.t_mayor, [2.0, 8.5, 4.2], primera_izq=True)

    _seccion(doc, f"Comparativo Regional — Occidente / Centro-Bajío ({c.act})")
    _p(doc, c.x_region, despues=5)
    _tabla(doc, ["Lugar", "Estado", "TD (%)", "vs Media (pp)"], c.t_region,
           [2.0, 5.5, 3.0, 4.2], primera_izq=True)

    _seccion(doc, "Indicadores Complementarios del Mercado Laboral",
             f"Comparativo anual: {c.act} vs {c.ant} — Jalisco")
    _p(doc, c.x_indicadores, despues=5)
    _tabla(doc, ["Indicador", c.ant, c.act, "Var. (pp)", f"Nacional {c.act}"],
           c.t_indicadores, [5.5, 2.2, 2.2, 2.2, 2.6], primera_izq=True)

    _seccion(doc, f"Indicadores por Sexo — Jalisco {c.act}")
    _p(doc, c.x_sexo, despues=5)
    _tabla(doc, ["Indicador", "Total", "Hombres", "Mujeres", "Brecha M−H"], c.t_sexo,
           [3.5, 2.6, 2.6, 2.6, 3.2], primera_izq=True)
    _p(doc, [(f"Variación anual por sexo ({c.ant} → {c.act}, pp):", 1)])
    _p(doc, c.x_sexo_var, antes=0)
    _tabla(doc, ["Indicador", "Hombres", "Mujeres"], c.t_sexo_var, [3.5, 3.0, 3.0], primera_izq=True)
    _p(doc, NOTA_SEXO, T_CHICA, italic=True, antes=0)

    if c.t_ocupacion:
        _seccion(doc, f"Población Ocupada — Jalisco {c.act}", "Personas y variación (absoluta y %)")
        n_comp = len(c.t_ocupacion[0]) - 2
        anchos = [4.6, 2.2] + [8.2 / n_comp] * n_comp
        _p(doc, c.x_ocupacion, despues=5)
        _tabla(doc, *c.t_ocupacion, anchos, primera_izq=True, negritas=c.grupos_ocupacion)
        if c.nota_sectores:
            _p(doc, c.nota_sectores, T_CHICA, italic=True, antes=0)

    doc.save(ruta)


# ── Uso directo ───────────────────────────────────────────────────────────────

def generar(periodo: str, carpeta_datos="datos", carpeta_salida="fichas") -> tuple[Path, Path]:
    anio, trim = map(int, re.fullmatch(r"(\d{4})_T(\d)", periodo).groups())
    previo = f"{anio - 1}_T4" if trim == 1 else f"{anio}_T{trim - 1}"
    c = preparar(calcular(periodo, carpeta_datos), calcular(f"{anio - 1}_T{trim}", carpeta_datos),
                 calcular(f"{anio - 1}_T4", carpeta_datos), calcular(previo, carpeta_datos))
    salida = Path(carpeta_salida)
    salida.mkdir(parents=True, exist_ok=True)
    txt, docx = salida / f"ficha_{periodo}.txt", salida / f"ficha_{periodo}.docx"
    generar_txt(c, txt)
    generar_docx(c, docx)
    return txt, docx


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description="Genera la ficha ENOE Jalisco (.txt y .docx)")
    ap.add_argument("--periodo", required=True, help="AAAA_T# (ej. 2026_T2)")
    args = ap.parse_args()
    for ruta in generar(args.periodo):
        print(f"  ✅ {ruta}")
